"""
engine.py – VerifAI  (multi-mode, streaming, OCR-enabled)
"""
from __future__ import annotations
import re, os, time
import fitz
import ollama

MODEL     = "llama3.2:3b"
MAX_CHARS = 4000

RISK_KEYWORDS = [
    "indemnif", "liability", "unlimited", "termination", "penalty",
    "liquidated damages", "arbitration", "governing law", "jurisdiction",
    "waive", "forfeit", "breach", "non-compete", "non-solicitation",
    "intellectual property", "assignment", "sublicense", "confidential",
    "auto-renew", "force majeure", "consequential damages",
    "sole discretion", "without cause",
]

_INLINE_RULES = """
STRICT FORMATTING RULES:
- Every bullet line starts with: •
- Every suggestion line starts with: →  and IMMEDIATELY follows its bullet (no blank line between)
- Page refs: use the [PAGE N] numbers from the document. Line refs: use L1, L2 etc.
- NEVER write "exact quote from document" — always copy the REAL text from the document
- NEVER invent page numbers — only cite [PAGE N] markers you can see in the document text
- If you cannot find real text to quote, write what the clause says in your own words"""

PROMPTS = {
    "contract": """You are a Senior Legal Auditor analyzing a real legal contract.
Reply in EXACTLY this structure — no intro, no sign-off:

EXECUTIVE SUMMARY:
[1-2 sentences: contract type, parties, purpose, value/date if stated]

KEY RISK FLAGS:
• RISK TYPE | Page N, Line N: "copy the actual clause text here" — why risky
→ Specific fix: what to negotiate or replace
[repeat for each risk found]

LIABILITY & INDEMNIFICATION:
• Page N, Line N: "copy the actual text" — financial exposure
→ How to limit or balance this
[repeat]

TERMINATION CONDITIONS:
• Page N, Line N: "copy the actual text" — trigger and consequence
→ Fairer language to request
[repeat]

AMBIGUOUS OR MISSING CLAUSES:
• Page N or Section name: describe what is vague or absent
→ Exact language to add or clarify
[repeat]

OVERALL RISK SCORE:
[LOW / MEDIUM / HIGH / CRITICAL] — one sentence reason

FINAL VERDICT:
[REJECT THIS CONTRACT / NEGOTIATE BEFORE SIGNING / ACCEPTABLE WITH CAUTION / SAFE TO SIGN]
1-2 plain English sentences the client should act on immediately.
""" + _INLINE_RULES,

    "nda": """You are a privacy law expert analyzing an NDA or Privacy Policy.
Reply in EXACTLY this structure:

DOCUMENT SUMMARY:
[1-2 sentences: type, parties, scope of confidentiality]

DISCLOSURE RISKS:
• Page N, Line N: "copy actual text" — risk explanation
→ Specific protection to request
[repeat]

DATA HANDLING CLAUSES:
• Page N, Line N: "copy actual text" — how data is handled
→ Protection language to add
[repeat]

DURATION & EXPIRY:
• Page N, Line N: "copy actual text" — when obligations end
→ Fairer duration to request
[repeat]

MISSING PROTECTIONS:
• Page N or Section name: what standard NDA protection is absent
→ Exact language to insert
[repeat]

OVERALL RISK SCORE:
[LOW / MEDIUM / HIGH / CRITICAL] — one sentence reason

FINAL VERDICT:
[REJECT THIS NDA / NEGOTIATE BEFORE SIGNING / ACCEPTABLE WITH CAUTION / SAFE TO SIGN]
1-2 plain English sentences on next steps.
""" + _INLINE_RULES,

    "employment": """You are an employment law expert analyzing an employment agreement.
Reply in EXACTLY this structure:

DOCUMENT SUMMARY:
[1-2 sentences: role, parties, compensation if stated]

COMPENSATION & BENEFITS FLAGS:
• Page N, Line N: "copy actual text" — unusual or missing term
→ What to negotiate
[repeat]

NON-COMPETE & RESTRICTIONS:
• Page N, Line N: "copy actual text" — restriction described
→ How to narrow or limit this
[repeat]

TERMINATION & SEVERANCE:
• Page N, Line N: "copy actual text" — how employment ends
→ Fairer severance language to request
[repeat]

IP & WORK OWNERSHIP:
• Page N, Line N: "copy actual text" — who owns work created
→ How to carve out personal projects or prior IP
[repeat]

MISSING PROTECTIONS:
• Page N or Section name: absent standard protection
→ Language to request added
[repeat]

OVERALL RISK SCORE:
[LOW / MEDIUM / HIGH / CRITICAL] — one sentence reason

FINAL VERDICT:
[DO NOT SIGN / NEGOTIATE BEFORE SIGNING / ACCEPTABLE WITH CAUTION / SAFE TO SIGN]
1-2 plain English sentences for the employee.
""" + _INLINE_RULES,

    "tos": """You are a consumer rights expert analyzing Terms of Service or a EULA.
Reply in EXACTLY this structure:

DOCUMENT SUMMARY:
[1-2 sentences: service, provider, key obligations]

USER RIGHTS RISKS:
• Page N, Line N: "copy actual text" — user right restricted
→ What to push back on or opt out of
[repeat]

DATA & PRIVACY CLAUSES:
• Page N, Line N: "copy actual text" — how user data is handled
→ Privacy controls to request
[repeat]

LIABILITY WAIVERS:
• Page N, Line N: "copy actual text" — what company disclaims
→ Counter-protection to seek
[repeat]

TERMINATION OF SERVICE:
• Page N, Line N: "copy actual text" — when service ends
→ Data export or notice period to negotiate
[repeat]

MISSING OR UNFAIR CLAUSES:
• Page N or Section name: what is absent or one-sided
→ Balancing language to request
[repeat]

OVERALL RISK SCORE:
[LOW / MEDIUM / HIGH / CRITICAL] — one sentence reason

FINAL VERDICT:
[AVOID THIS SERVICE / USE WITH CAUTION / ACCEPTABLE / SAFE TO USE]
1-2 plain English sentences on whether to use this service.
""" + _INLINE_RULES,

    "compare": """You are a Senior Legal Auditor comparing two contracts.
Reply in EXACTLY this structure:

COMPARISON SUMMARY:
[2-3 sentences: what the documents are and key differences at a glance]

FAVOURABLE DIFFERENCES (Document B vs A):
• Topic: Doc A says "quote" vs Doc B says "quote" — why B is better
→ Whether to adopt as-is or improve further
[repeat]

UNFAVOURABLE DIFFERENCES (Document B vs A):
• Topic: Doc A says "quote" vs Doc B says "quote" — why B is worse
→ Change to bring Document B up to standard
[repeat]

CLAUSES ONLY IN DOCUMENT A:
• Page N: clause absent in Document B
→ Whether Document B should adopt this
[repeat]

CLAUSES ONLY IN DOCUMENT B:
• Page N: clause absent in Document A
→ Whether Document A should adopt this
[repeat]

OVERALL RISK SCORE:
Document A: [LOW/MEDIUM/HIGH/CRITICAL] — Document B: [LOW/MEDIUM/HIGH/CRITICAL]

FINAL VERDICT:
[PREFER DOCUMENT A / PREFER DOCUMENT B / NEITHER IS ACCEPTABLE / BOTH ACCEPTABLE]
1-2 plain English sentences on which to sign and first negotiation priority.
""" + _INLINE_RULES,
}

MODE_LABELS = {
    "contract":   "Contract Audit",
    "nda":        "NDA / Privacy Policy",
    "employment": "Employment Agreement",
    "tos":        "Terms of Service / EULA",
    "compare":    "Document Comparison",
}

# ── model utils ───────────────────────────────────────────────────────────────

def get_available_models() -> list[str]:
    try:
        return [m.model for m in ollama.list().models]
    except Exception:
        return []

def is_model_available(model_name: str) -> bool:
    return any(model_name in m or m in model_name
               for m in get_available_models())

# ── OCR helper ────────────────────────────────────────────────────────────────

def _ocr_pdf(path: str) -> str:
    """Use pytesseract OCR to extract text from scanned PDF pages."""
    try:
        import pytesseract
        from PIL import Image
        import fitz as _fitz

        parts = []
        with _fitz.open(path) as doc:
            for pg_num, page in enumerate(doc, 1):
                # Render at 200 DPI for good OCR accuracy
                mat  = _fitz.Matrix(200 / 72, 200 / 72)
                pix  = page.get_pixmap(matrix=mat, colorspace=_fitz.csGRAY)
                img  = Image.frombytes("L", [pix.width, pix.height], pix.samples)
                text = pytesseract.image_to_string(img, lang="eng")
                lines = [l.strip() for l in text.splitlines() if l.strip()]
                if lines:
                    numbered = "\n".join(f"L{i}: {t}" for i, t in enumerate(lines, 1))
                    parts.append(f"[PAGE {pg_num}]\n{numbered}")

        return "\n\n".join(parts) if parts else ""

    except ImportError:
        return ""   # pytesseract not available
    except Exception:
        return ""

# ── text extraction ───────────────────────────────────────────────────────────

def extract_text(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    if   ext == ".pdf":  return _extract_pdf(file_path)
    elif ext == ".docx": return _extract_docx(file_path)
    elif ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
    else:
        raise ValueError(f"Unsupported file type '{ext}'. Use PDF, DOCX, or TXT.")

def _extract_pdf(path: str) -> str:
    """Extract text from PDF. Falls back to OCR for scanned images."""
    parts = []
    with fitz.open(path) as doc:
        for pg, page in enumerate(doc, 1):
            raw_lines = [l.strip() for l in
                         page.get_text("text").splitlines() if l.strip()]
            if raw_lines:
                numbered = "\n".join(f"L{i}: {t}"
                                     for i, t in enumerate(raw_lines, 1))
                parts.append(f"[PAGE {pg}]\n{numbered}")

    if parts:
        # Check if we got meaningful text (not just noise)
        total_chars = sum(len(p) for p in parts)
        if total_chars > 100:
            return "\n\n".join(parts)

    # Fallback: OCR for scanned PDFs
    ocr_text = _ocr_pdf(path)
    if ocr_text:
        return ocr_text

    raise ValueError(
        "This PDF appears to be a scanned image with no extractable text.\n"
        "OCR was attempted but failed. Make sure pytesseract is installed:\n"
        "  pip install pytesseract\n"
        "  brew install tesseract  (macOS)"
    )

def _extract_docx(path: str) -> str:
    try:
        import docx
        from docx.oxml.ns import qn
        doc  = docx.Document(path)
        out  = []
        pg   = 1
        ln   = 1
        out.append(f"[PAGE {pg}]")

        for para in doc.paragraphs:
            xml = para._element
            if (xml.findall(f'.//{qn("w:br")}[@{qn("w:type")}="page"]') or
                    xml.findall(f'.//{qn("w:lastRenderedPageBreak")}')):
                pg += 1; ln = 1
                out.append(f"\n[PAGE {pg}]")
            text = para.text.strip()
            if text:
                out.append(f"L{ln}: {text}")
                ln += 1

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    t = cell.text.strip()
                    if t:
                        out.append(f"L{ln}: {t}")
                        ln += 1

        result = "\n".join(out)

        # Fallback: character-count page estimate if no real breaks found
        if pg == 1:
            out2 = ["[PAGE 1]"]
            pg2 = 1; ln2 = 1; cc = 0
            for para in doc.paragraphs:
                t = para.text.strip()
                if not t:
                    continue
                cc += len(t)
                if cc > 2500:
                    pg2 += 1; ln2 = 1; cc = 0
                    out2.append(f"\n[PAGE {pg2}]")
                out2.append(f"L{ln2}: {t}")
                ln2 += 1
            result = "\n".join(out2)

        return result
    except ImportError:
        raise RuntimeError("Run: pip install python-docx")

def _clean(text: str) -> str:
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    text = re.sub(r"[^\x09\x0A\x20-\x7E\u00A0-\uFFFF]", " ", text)
    return text.strip()

def _pre_scan(text: str) -> list[str]:
    lower = text.lower()
    return [kw for kw in RISK_KEYWORDS if kw in lower]

# ── streaming ─────────────────────────────────────────────────────────────────

def audit_contract_stream(file_path: str, mode: str = "contract"):
    """Yields text chunks as the model generates them."""
    raw   = extract_text(file_path)
    text  = _clean(raw)

    if len(text) < 50:
        raise ValueError(
            "Could not extract readable text from this file.\n"
            "For scanned PDFs, ensure tesseract is installed:\n"
            "  brew install tesseract"
        )

    risks = _pre_scan(text)

    risk_note = ""
    if risks and mode != "compare":
        risk_note = (
            f"\n\nPRE-SCAN NOTE — these risk terms appear in the document. "
            f"Find each one and cite its exact [PAGE N] L<n> location: "
            f"{', '.join(risks)}"
        )

    system   = PROMPTS.get(mode, PROMPTS["contract"])
    user_msg = (
        f"Analyze this document thoroughly. "
        f"Copy REAL text from the document for every quote — "
        f"NEVER write placeholder text like 'exact quote from document'."
        f"{risk_note}\n\n"
        f"DOCUMENT TEXT (use [PAGE N] and L<n> markers for all citations):\n"
        f"{text[:MAX_CHARS]}"
    )

    try:
        stream = ollama.chat(
            model=MODEL,
            messages=[
                {"role": "system", "content": system},
                {"role": "user",   "content": user_msg},
            ],
            stream=True,
            options={
                "temperature": 0.05,
                "num_predict": 1800,
                "num_ctx":     4096,
            },
        )
        for chunk in stream:
            delta = chunk["message"]["content"]
            if delta:
                yield delta

    except ollama.ResponseError as e:
        raise RuntimeError(
            f"Model '{MODEL}' error: {e.error}\n"
            f"Run: ollama pull {MODEL}"
        ) from e
    except Exception as e:
        raise RuntimeError(
            f"Cannot reach Ollama. Start with: ollama serve\n\n{e}"
        ) from e


def compare_stream(file_a: str, file_b: str):
    text_a = _clean(extract_text(file_a))[:MAX_CHARS // 2]
    text_b = _clean(extract_text(file_b))[:MAX_CHARS // 2]

    if len(text_a) < 50 or len(text_b) < 50:
        raise ValueError("One or both files have no extractable text.")

    user_msg = (
        f"Compare these two documents. Copy REAL text for every quote.\n\n"
        f"=== DOCUMENT A ===\n{text_a}\n\n"
        f"=== DOCUMENT B ===\n{text_b}"
    )
    try:
        stream = ollama.chat(
            model=MODEL,
            messages=[
                {"role": "system", "content": PROMPTS["compare"]},
                {"role": "user",   "content": user_msg},
            ],
            stream=True,
            options={"temperature": 0.05, "num_predict": 1800, "num_ctx": 4096},
        )
        for chunk in stream:
            delta = chunk["message"]["content"]
            if delta:
                yield delta
    except Exception as e:
        raise RuntimeError(f"Compare failed: {e}") from e


def audit_contract(file_path: str, mode: str = "contract") -> str:
    return "".join(audit_contract_stream(file_path, mode))